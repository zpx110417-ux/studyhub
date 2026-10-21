"""管理后台：审批 / 编辑授权 / 举报处置 / 封禁 / 申诉审批 / 用户管理 / 页面删除 / 站点设置。"""
from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    redirect,
    render_template,
    request,
    url_for,
)

from config import ADMIN_MAX_COUNT, BAN_DAY_CHOICES
from .db import get_db
from .helpers import admin_required, current_user, get_setting, set_setting

bp = Blueprint("admin", __name__, url_prefix="/admin")


# ---------- 概览 ----------

@bp.route("/")
@admin_required
def index():
    db = get_db()
    stats = {
        "pending_pages": db.execute(
            "SELECT COUNT(*) FROM pages WHERE status='pending'").fetchone()[0],
        "pending_requests": db.execute(
            "SELECT COUNT(*) FROM edit_requests WHERE status='pending'").fetchone()[0],
        "pending_reports": db.execute(
            "SELECT COUNT(*) FROM reports WHERE status='pending'").fetchone()[0],
        "pending_appeals": db.execute(
            "SELECT COUNT(*) FROM appeals WHERE status='pending'").fetchone()[0],
        "active_bans": db.execute(
            "SELECT COUNT(*) FROM bans WHERE lifted_at IS NULL "
            "AND banned_until > datetime('now','localtime')").fetchone()[0],
        "pending_suggestions": db.execute(
            "SELECT COUNT(*) FROM suggestions WHERE handled=0").fetchone()[0],
    }
    from .helpers import online_users

    online = online_users()
    return render_template("admin/index.html", stats=stats, online=online)


# ---------- 页面审批 ----------

@bp.route("/review")
@admin_required
def review():
    rows = get_db().execute(
        "SELECT p.*, u.username AS author, r.username AS reviewer FROM pages p "
        "JOIN users u ON u.id=p.author_id LEFT JOIN users r ON r.id=p.reviewed_by "
        "WHERE p.status='pending' ORDER BY p.updated_at"
    ).fetchall()
    return render_template("admin/review.html", pages=rows)


@bp.route("/review/<int:page_id>/approve", methods=("POST",))
@admin_required
def approve_page(page_id):
    db = get_db()
    if db.execute(
        "SELECT 1 FROM pages WHERE id=? AND status='pending'", (page_id,)
    ).fetchone() is None:
        abort(404)
    db.execute(
        "UPDATE pages SET status='published', reviewed_by=?, "
        "reviewed_at=datetime('now','localtime'), review_note='', reject_count=0 WHERE id=?",
        (current_user()["id"], page_id),
    )
    db.commit()
    flash("页面已批准发布。", "success")
    return redirect(url_for("admin.review"))


@bp.route("/review/<int:page_id>/reject", methods=("POST",))
@admin_required
def reject_page(page_id):
    note = request.form.get("note", "").strip()
    if not note:
        flash("请填写驳回理由，便于作者修改。", "danger")
        return redirect(url_for("admin.review"))
    db = get_db()
    if db.execute(
        "SELECT 1 FROM pages WHERE id=? AND status='pending'", (page_id,)
    ).fetchone() is None:
        abort(404)
    db.execute(
        "UPDATE pages SET status='rejected', reject_count=reject_count+1, "
        "last_reject_date=date('now','localtime'), review_note=?, reviewed_by=?, "
        "reviewed_at=datetime('now','localtime') WHERE id=?",
        (note, current_user()["id"], page_id),
    )
    db.commit()
    flash("已驳回，作者可修改后重新提交（连续驳回 3 次当天不可再提交）。", "success")
    return redirect(url_for("admin.review"))


# ---------- 编辑授权申请 ----------

@bp.route("/requests")
@admin_required
def requests():
    rows = get_db().execute(
        "SELECT er.*, u.username AS requester, r.username AS reviewer, p.title AS page_title "
        "FROM edit_requests er JOIN users u ON u.id=er.requester_id "
        "LEFT JOIN users r ON r.id=er.reviewed_by JOIN pages p ON p.id=er.page_id "
        "ORDER BY CASE er.status WHEN 'pending' THEN 0 ELSE 1 END, er.id DESC"
    ).fetchall()
    return render_template("admin/requests.html", rows=rows)


@bp.route("/requests/<int:req_id>/<action>", methods=("POST",))
@admin_required
def handle_request(req_id, action):
    if action not in ("approve", "reject", "revoke"):
        abort(404)
    db = get_db()
    req = db.execute("SELECT * FROM edit_requests WHERE id=?", (req_id,)).fetchone()
    if req is None:
        abort(404)
    hours = current_app.config["EDIT_GRANT_HOURS"]
    if action == "approve" and req["status"] == "pending":
        db.execute(
            "UPDATE edit_requests SET status='approved', reviewed_by=?, "
            f"expires_at=datetime('now','localtime', '+{hours} hours') WHERE id=?",
            (current_user()["id"], req_id),
        )
        flash(f"已授权：该用户可在 {hours} 小时内编辑该页面。", "success")
    elif action == "reject" and req["status"] == "pending":
        db.execute(
            "UPDATE edit_requests SET status='rejected', reviewed_by=?, "
            "reviewed_at=datetime('now','localtime') WHERE id=?",
            (current_user()["id"], req_id),
        )
        flash("已驳回该编辑申请。", "success")
    elif action == "revoke" and req["status"] == "approved":
        db.execute("UPDATE edit_requests SET status='revoked' WHERE id=?", (req_id,))
        flash("已撤销该授权。", "success")
    else:
        flash("该申请不在可操作状态。", "warning")
    db.commit()
    return redirect(url_for("admin.requests"))


# ---------- 举报处置 ----------

@bp.route("/reports")
@admin_required
def reports():
    rows = get_db().execute(
        "SELECT rep.*, ru.username AS reporter, tu.username AS target, "
        "p.title AS page_title, p.id AS pid, "
        "q.title AS question_title, q.id AS qid, a.question_id AS aqid, "
        "pc.post_id AS pc_post_id "
        "FROM reports rep "
        "LEFT JOIN users ru ON ru.id=rep.reporter_id "
        "JOIN users tu ON tu.id=rep.target_user_id "
        "LEFT JOIN pages p ON p.id=rep.page_id "
        "LEFT JOIN questions q ON q.id=rep.question_id "
        "LEFT JOIN answers a ON a.id=rep.answer_id "
        "LEFT JOIN posts po ON po.id=rep.post_id "
        "LEFT JOIN post_comments pc ON pc.id=rep.post_comment_id "
        "ORDER BY CASE rep.status WHEN 'pending' THEN 0 ELSE 1 END, rep.id DESC"
    ).fetchall()
    return render_template("admin/reports.html", rows=rows, ban_days=BAN_DAY_CHOICES)


@bp.route("/reports/<int:report_id>/dismiss", methods=("POST",))
@admin_required
def dismiss_report(report_id):
    db = get_db()
    db.execute(
        "UPDATE reports SET status='dismissed', handled_by=?, "
        "handled_at=datetime('now','localtime') WHERE id=? AND status='pending'",
        (current_user()["id"], report_id),
    )
    db.commit()
    flash("已驳回该举报。", "success")
    return redirect(url_for("admin.reports"))


def _report_evidence(db, report_id):
    """抓取被举报内容的原文快照，作为封禁依据存入封禁记录，供申诉审批时参考。"""
    rep = db.execute(
        "SELECT page_id, question_id, answer_id, post_id, post_comment_id FROM reports WHERE id=?",
        (report_id,),
    ).fetchone()
    if rep is None:
        return ""
    if rep["page_id"]:
        row = db.execute(
            "SELECT content_md FROM pages WHERE id=?", (rep["page_id"],)
        ).fetchone()
    elif rep["question_id"]:
        row = db.execute(
            "SELECT content_md FROM questions WHERE id=?", (rep["question_id"],)
        ).fetchone()
    elif rep["answer_id"]:
        row = db.execute(
            "SELECT content_md FROM answers WHERE id=?", (rep["answer_id"],)
        ).fetchone()
    elif rep["post_id"]:
        row = db.execute(
            "SELECT content_md FROM posts WHERE id=?", (rep["post_id"],)
        ).fetchone()
    elif rep["post_comment_id"]:
        row = db.execute(
            "SELECT content_md FROM post_comments WHERE id=?", (rep["post_comment_id"],)
        ).fetchone()
    else:
        row = None
    return (row["content_md"] if row else "")[:2000]


def _do_ban(report_id=None):
    """从举报处置页或用户管理页发起封禁。"""
    user_id = request.form.get("user_id", type=int)
    days = request.form.get("days", type=int)
    reason = request.form.get("reason", "").strip()
    if days not in BAN_DAY_CHOICES:
        flash("封禁时长不合法。", "danger")
        return redirect(url_for("admin.reports"))
    if not reason:
        flash("请填写封禁理由。", "danger")
        return redirect(url_for("admin.reports"))
    db = get_db()
    target = db.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if target is None:
        abort(404)
    if target["role"] == "admin":
        flash("不能封禁管理员账号。", "danger")
        return redirect(url_for("admin.reports"))
    active = db.execute(
        "SELECT 1 FROM bans WHERE user_id=? AND lifted_at IS NULL "
        "AND banned_until > datetime('now','localtime')",
        (user_id,),
    ).fetchone()
    if active:
        flash("该用户已在封禁中。", "warning")
        return redirect(url_for("admin.reports"))
    db.execute(
        "INSERT INTO bans (user_id, reason, evidence, banned_by, banned_until) "
        "VALUES (?, ?, ?, ?, datetime('now','localtime', ?))",
        (
            user_id, reason,
            _report_evidence(db, report_id) if report_id else "",
            current_user()["id"], f"+{days} days",
        ),
    )
    if report_id:
        db.execute(
            "UPDATE reports SET status='actioned', handled_by=?, "
            "handled_at=datetime('now','localtime') WHERE id=? AND status='pending'",
            (current_user()["id"], report_id),
        )
    db.commit()
    flash(f"已将 {target['username']} 封禁 {days} 天。", "success")
    return redirect(url_for("admin.reports"))


@bp.route("/reports/<int:report_id>/ban", methods=("POST",))
@admin_required
def ban_from_report(report_id):
    return _do_ban(report_id=report_id)


# ---------- 用户管理 ----------

@bp.route("/users")
@admin_required
def users():
    rows = get_db().execute(
        "SELECT u.*, (SELECT b.banned_until FROM bans b WHERE b.user_id=u.id "
        "AND b.lifted_at IS NULL AND b.banned_until > datetime('now','localtime') "
        "ORDER BY b.banned_until DESC LIMIT 1) AS banned_until "
        "FROM users u ORDER BY u.id"
    ).fetchall()
    return render_template("admin/users.html", rows=rows, ban_days=BAN_DAY_CHOICES)


@bp.route("/users/<int:user_id>/role", methods=("POST",))
@admin_required
def set_role(user_id):
    role = request.form.get("role")
    if role not in ("user", "admin"):
        abort(400)
    if user_id == current_user()["id"]:
        flash("不能修改自己的角色。", "danger")
        return redirect(url_for("admin.users"))
    db = get_db()
    target = db.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if target is None:
        abort(404)
    if role == "admin" and target["role"] != "admin":
        count = db.execute("SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0]
        if count >= ADMIN_MAX_COUNT:
            flash(f"管理员数量已达上限（{ADMIN_MAX_COUNT} 位）。", "danger")
            return redirect(url_for("admin.users"))
    db.execute("UPDATE users SET role=? WHERE id=?", (role, user_id))
    db.commit()
    flash("角色已更新。", "success")
    return redirect(url_for("admin.users"))


@bp.route("/users/<int:user_id>/ban", methods=("POST",))
@admin_required
def ban_user(user_id):
    if request.form.get("user_id", type=int) != user_id:
        abort(400)
    return _do_ban()


@bp.route("/users/import", methods=("POST",))
@admin_required
def import_users():
    """导入成员名单 Excel：第一列用户名，第二列密码，一一对应批量创建普通用户。"""
    file = request.files.get("file")
    if file is None or not file.filename:
        flash("请选择要导入的 Excel 文件。", "danger")
        return redirect(url_for("admin.users"))
    if not file.filename.lower().endswith(".xlsx"):
        flash("仅支持 .xlsx 格式的 Excel 文件（在 Excel/WPS 中另存为 .xlsx）。", "danger")
        return redirect(url_for("admin.users"))

    from openpyxl import load_workbook
    from werkzeug.security import generate_password_hash
    from .db import USERNAME_RE

    ok, skipped = [], []
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
            username = str(row[0]).strip() if row and row[0] is not None else ""
            password = str(row[1]).strip() if row and len(row) > 1 and row[1] is not None else ""
            if not username and not password:
                continue  # 空行直接忽略
            if not USERNAME_RE.fullmatch(username):
                skipped.append(f"第{row_no}行 用户名不合规（需3~30位字母/数字/下划线/中文）")
                continue
            if len(password) < 6:
                skipped.append(f"第{row_no}行（{username}）密码不足 6 位")
                continue
            if get_db().execute(
                "SELECT 1 FROM users WHERE username=?", (username,)
            ).fetchone():
                skipped.append(f"第{row_no}行 用户名已存在：{username}")
                continue
            get_db().execute(
                "INSERT INTO users (username, password_hash) VALUES (?, ?)",
                (username, generate_password_hash(password, method="pbkdf2:sha256")),
            )
            ok.append(username)
        get_db().commit()
        wb.close()
    except Exception:
        current_app.logger.exception("成员名单导入失败")
        flash("文件解析失败，请确认是有效的 .xlsx 文件后重试。", "danger")
        return redirect(url_for("admin.users"))

    if ok:
        flash(f"导入成功 {len(ok)} 个账号：{'、'.join(ok)}", "success")
    else:
        flash("没有创建任何账号。", "warning")
    if skipped:
        preview = "；".join(skipped[:8]) + ("……" if len(skipped) > 8 else "")
        flash(f"跳过 {len(skipped)} 行：{preview}", "warning")
    return redirect(url_for("admin.users"))


# ---------- 封禁（小黑屋）管理 ----------

@bp.route("/bans")
@admin_required
def bans():
    rows = get_db().execute(
        "SELECT b.*, u.username AS banned_user, a.username AS admin_name "
        "FROM bans b JOIN users u ON u.id=b.user_id JOIN users a ON a.id=b.banned_by "
        "ORDER BY (b.lifted_at IS NULL AND b.banned_until > datetime('now','localtime')) DESC, b.id DESC"
    ).fetchall()
    now_str = get_db().execute(
        "SELECT datetime('now','localtime') AS n"
    ).fetchone()["n"]
    return render_template("admin/bans.html", rows=rows, now_str=now_str)


@bp.route("/bans/<int:ban_id>/lift", methods=("POST",))
@admin_required
def lift_ban(ban_id):
    db = get_db()
    db.execute(
        "UPDATE bans SET lifted_at=datetime('now','localtime') "
        "WHERE id=? AND lifted_at IS NULL",
        (ban_id,),
    )
    db.commit()
    flash("已提前解封。", "success")
    return redirect(url_for("admin.bans"))


# ---------- 申诉审批（每条封禁最多 5 次，含执行封禁的管理员；每人限审 1 次） ----------

@bp.route("/appeals")
@admin_required
def appeals():
    db = get_db()
    me = current_user()["id"]
    rows = db.execute(
        "SELECT ap.*, au.username AS appellant, ra.username AS reviewer, "
        "b.reason AS ban_reason, b.evidence, b.banned_until, b.banned_by, "
        "ba.username AS banner "
        "FROM appeals ap JOIN bans b ON b.id=ap.ban_id "
        "JOIN users au ON au.id=ap.appellant_id "
        "JOIN users ba ON ba.id=b.banned_by "
        "LEFT JOIN users ra ON ra.id=ap.reviewed_by "
        "ORDER BY CASE ap.status WHEN 'pending' THEN 0 ELSE 1 END, ap.id DESC"
    ).fetchall()
    items = []
    for r in rows:
        reviewed_by_me = False
        if r["status"] == "pending":
            reviewed_by_me = db.execute(
                "SELECT COUNT(*) FROM appeals WHERE ban_id=? AND reviewed_by=?",
                (r["ban_id"], me),
            ).fetchone()[0] > 0
        items.append({
            "row": r,
            "can_review": r["status"] == "pending" and not reviewed_by_me,
        })
    return render_template("admin/appeals.html", items=items)


@bp.route("/appeals/<int:appeal_id>/<action>", methods=("POST",))
@admin_required
def handle_appeal(appeal_id, action):
    if action not in ("approve", "reject"):
        abort(404)
    db = get_db()
    ap = db.execute(
        "SELECT * FROM appeals WHERE id=? AND status='pending'", (appeal_id,)
    ).fetchone()
    if ap is None:
        abort(404)
    ban = db.execute("SELECT * FROM bans WHERE id=?", (ap["ban_id"],)).fetchone()
    me = current_user()["id"]
    # 每条封禁最多 5 次申诉审批，执行封禁的管理员也在其中，但每人限审一次
    if db.execute(
        "SELECT COUNT(*) FROM appeals WHERE ban_id=? AND reviewed_by=?",
        (ban["id"], me),
    ).fetchone()[0]:
        abort(403, "每条封禁的申诉，每位管理员只能审批一次。")
    if action == "approve":
        db.execute(
            "UPDATE appeals SET status='approved', reviewed_by=?, "
            "reviewed_at=datetime('now','localtime') WHERE id=?",
            (me, appeal_id),
        )
        db.execute(
            "UPDATE bans SET lifted_at=datetime('now','localtime') "
            "WHERE id=? AND lifted_at IS NULL",
            (ban["id"],),
        )
        flash("申诉成立，已解封该用户。", "success")
    else:
        db.execute(
            "UPDATE appeals SET status='rejected', reviewed_by=?, "
            "reviewed_at=datetime('now','localtime') WHERE id=?",
            (me, appeal_id),
        )
        flash("已驳回申诉，维持封禁。", "success")
    db.commit()
    return redirect(url_for("admin.appeals"))


# ---------- 页面管理 ----------

@bp.route("/pages")
@admin_required
def pages():
    rows = get_db().execute(
        "SELECT p.*, u.username AS author FROM pages p "
        "JOIN users u ON u.id=p.author_id ORDER BY p.updated_at DESC"
    ).fetchall()
    return render_template("admin/pages.html", pages=rows)


def _record_deletion(db, target_type, title, reason):
    """删除公示：文章名、理由、删除者公开存档。"""
    db.execute(
        "INSERT INTO deletions (target_type, title, reason, admin_id, admin_name) "
        "VALUES (?, ?, ?, ?, ?)",
        (target_type, title, reason, current_user()["id"], current_user()["username"]),
    )


@bp.route("/pages/<int:page_id>/delete", methods=("POST",))
@admin_required
def delete_page(page_id):
    reason = request.form.get("reason", "").strip()
    if not reason:
        flash("删除必须给出理由，理由将公开发布到删除公示。", "danger")
        return redirect(url_for("admin.pages"))
    db = get_db()
    row = db.execute("SELECT title FROM pages WHERE id=?", (page_id,)).fetchone()
    if row is None:
        abort(404)
    _record_deletion(db, "page", row["title"], reason)
    db.execute("DELETE FROM pages WHERE id=?", (page_id,))  # 版本/申请记录级联删除
    db.commit()
    flash(f"页面《{row['title']}》已删除，删除信息已公示。", "success")
    return redirect(url_for("admin.pages"))


@bp.route("/questions/<int:question_id>/delete", methods=("POST",))
@admin_required
def delete_question(question_id):
    reason = request.form.get("reason", "").strip()
    if not reason:
        flash("删除必须给出理由，理由将公开发布到删除公示。", "danger")
        return redirect(url_for("qa.list_questions"))
    db = get_db()
    row = db.execute("SELECT title FROM questions WHERE id=?", (question_id,)).fetchone()
    if row is None:
        abort(404)
    _record_deletion(db, "question", row["title"], reason)
    db.execute("DELETE FROM questions WHERE id=?", (question_id,))  # 回答级联删除
    db.commit()
    flash(f"问答帖《{row['title']}》已删除，删除信息已公示。", "success")
    return redirect(url_for("qa.list_questions"))


@bp.route("/posts/<int:post_id>/delete", methods=("POST",))
@admin_required
def delete_post(post_id):
    reason = request.form.get("reason", "").strip()
    if not reason:
        flash("删除必须给出理由，理由将公开发布到删除公示。", "danger")
        return redirect(url_for("newspaper.index"))
    db = get_db()
    row = db.execute("SELECT content_md FROM posts WHERE id=?", (post_id,)).fetchone()
    if row is None:
        abort(404)
    title = (row["content_md"].splitlines()[0] if row["content_md"] else "")[:30] or "（无文字内容）"
    _record_deletion(db, "post", f"报刊动态：{title}…", reason)
    db.execute("DELETE FROM posts WHERE id=?", (post_id,))  # 点赞与评论级联删除
    db.commit()
    flash("报刊动态已删除，删除信息已公示。", "success")
    return redirect(url_for("newspaper.index"))


@bp.route("/post-comments/<int:comment_id>/delete", methods=("POST",))
@admin_required
def delete_post_comment(comment_id):
    reason = request.form.get("reason", "").strip()
    if not reason:
        flash("删除必须给出理由，理由将公开发布到删除公示。", "danger")
        return redirect(url_for("newspaper.index"))
    db = get_db()
    row = db.execute(
        "SELECT c.content_md, p.id AS pid, p.content_md AS post_md "
        "FROM post_comments c JOIN posts p ON p.id=c.post_id WHERE c.id=?",
        (comment_id,),
    ).fetchone()
    if row is None:
        abort(404)
    post_title = (row["post_md"].splitlines()[0] if row["post_md"] else "")[:20] or "无标题动态"
    title = (row["content_md"].splitlines()[0] if row["content_md"] else "")[:30]
    _record_deletion(
        db, "post_comment", f"报刊评论（对动态《{post_title}…》的评论）：{title}…", reason,
    )
    db.execute("DELETE FROM post_comments WHERE id=?", (comment_id,))
    db.commit()
    flash("评论已删除，删除信息已公示。", "success")
    return redirect(url_for("newspaper.index"))


# ---------- 意见箱（匿名，全部管理员共见） ----------

@bp.route("/suggestions")
@admin_required
def suggestions():
    rows = get_db().execute(
        "SELECT * FROM suggestions ORDER BY handled ASC, id DESC"
    ).fetchall()
    return render_template("admin/suggestions.html", rows=rows)


@bp.route("/suggestions/<int:sid>/handle", methods=("POST",))
@admin_required
def handle_suggestion(sid):
    db = get_db()
    db.execute("UPDATE suggestions SET handled=1 WHERE id=? AND handled=0", (sid,))
    db.commit()
    flash("意见已标记为已处理。", "success")
    return redirect(url_for("admin.suggestions"))


@bp.route("/suggestions/<int:sid>/delete", methods=("POST",))
@admin_required
def delete_suggestion(sid):
    db = get_db()
    db.execute("DELETE FROM suggestions WHERE id=?", (sid,))
    db.commit()
    flash("意见已删除。", "success")
    return redirect(url_for("admin.suggestions"))


# ---------- 站点设置（使用规则 / 敏感词词库） ----------

@bp.route("/settings", methods=("GET", "POST"))
@admin_required
def settings():
    if request.method == "POST":
        set_setting("rules_md", request.form.get("rules_md", ""))
        set_setting("words_high", request.form.get("words_high", ""))
        set_setting("words_low", request.form.get("words_low", ""))
        get_db().commit()
        flash("设置已保存。", "success")
        return redirect(url_for("admin.settings"))
    data = {k: get_setting(k, "") for k in ("rules_md", "words_high", "words_low")}
    last_scan = get_setting("last_scan_date", "")
    return render_template("admin/settings.html", data=data, last_scan=last_scan)
